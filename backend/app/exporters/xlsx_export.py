"""Excel (.xlsx) exporter via openpyxl."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.exporters.common import CSV_HEADERS, review_row
from app.models import Review


def export_xlsx(reviews: list[Review]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "reviews"
    ws.append(CSV_HEADERS)
    for review in reviews:
        row = review_row(review)
        ws.append([row[h] for h in CSV_HEADERS])
    for idx, header in enumerate(CSV_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(
            28, max(12, len(header) + 2)
        )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
